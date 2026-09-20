import io
import pytest
import fitz
import docx
from app.services.extractor import resume_extractor
from app.services.parser import resume_parser
from app.models.candidate import CandidateProfileSchema, CandidateBasics, ExperienceItem

def test_candidate_schema_nullability():
    """Verify that candidate profile schema does not invent fields and defaults cleanly."""
    profile = CandidateProfileSchema()
    assert profile.basics.name is None
    assert profile.basics.email is None
    assert profile.skills == []
    assert profile.experience == []
    assert profile.education == []
    assert profile.projects == []
    assert profile.certifications == []
    assert profile.achievements == []

def test_pdf_extraction_in_memory():
    """Create a minimal PDF in memory and verify PyMuPDF extracts text correctly."""
    doc = fitz.open()
    page = doc.new_page()
    page.insert_text((50, 72), "Jane Doe\njane.doe@example.com\nSoftware Engineer with Python and React.")
    pdf_bytes = doc.write()
    doc.close()

    extracted = resume_extractor.extract_text(pdf_bytes, "pdf")
    assert "Jane Doe" in extracted
    assert "jane.doe@example.com" in extracted
    assert "Python" in extracted

def test_docx_extraction_in_memory():
    """Create a minimal DOCX in memory and verify python-docx extracts text correctly."""
    doc = docx.Document()
    doc.add_heading("John Smith", level=1)
    doc.add_paragraph("john.smith@example.com | (555) 123-4567")
    doc.add_paragraph("Full Stack Developer specializing in TypeScript and Node.js.")
    
    stream = io.BytesIO()
    doc.save(stream)
    docx_bytes = stream.getvalue()

    extracted = resume_extractor.extract_text(docx_bytes, "docx")
    assert "John Smith" in extracted
    assert "john.smith@example.com" in extracted
    assert "TypeScript" in extracted

def test_heuristic_extraction_never_invents():
    """Verify factual fallback parser extracts only explicit patterns and does not hallucinate."""
    sample_text = """
    Alice Wonderland
    alice@wonderland.dev
    https://github.com/alicew
    https://linkedin.com/in/alicewonderland

    Skills:
    Python, Docker, PostgreSQL
    """
    profile = resume_parser.heuristic_extraction(sample_text)
    assert profile.basics.email == "alice@wonderland.dev"
    assert profile.basics.github == "https://github.com/alicew"
    assert profile.basics.linkedin == "https://linkedin.com/in/alicewonderland"
    assert "Python" in profile.skills
    assert "Docker" in profile.skills
    assert "PostgreSQL" in profile.skills
    # Unmentioned fields must remain empty/null!
    assert profile.experience == []
    assert profile.education == []
    assert profile.basics.phone is None
