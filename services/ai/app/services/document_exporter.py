import io
import fitz
import docx
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from typing import Dict, Any

class DocumentExporter:
    @staticmethod
    def export_resume_pdf(resume_data: Dict[str, Any]) -> bytes:
        """Generate an ATS-friendly, clean vector PDF using PyMuPDF (fitz)."""
        doc = fitz.open()
        # Standard letter size: 612 x 792 pt
        page = doc.new_page(width=612, height=792)

        margin_left = 54.0
        margin_right = 558.0
        max_width = margin_right - margin_left
        y = 54.0

        header = resume_data.get("header") or {}
        name = header.get("name") or "Candidate Name"
        contact_items = [
            header.get("email"),
            header.get("phone"),
            header.get("location"),
            header.get("linkedin"),
            header.get("github")
        ]
        contact_line = " | ".join([c for c in contact_items if c])

        # 1. Header (Name + Contact info)
        page.insert_text((margin_left, y), name, fontsize=18, fontname="helv", color=(0.1, 0.1, 0.15))
        y += 16
        if contact_line:
            page.insert_text((margin_left, y), contact_line, fontsize=9, fontname="helv", color=(0.3, 0.35, 0.4))
            y += 14

        # Divider line
        page.draw_line((margin_left, y), (margin_right, y), color=(0.7, 0.75, 0.8), width=0.75)
        y += 16

        def check_page_break(doc, page, current_y, needed_space=30):
            if current_y + needed_space > 740:
                new_p = doc.new_page(width=612, height=792)
                return new_p, 54.0
            return page, current_y

        def insert_section_heading(doc, page, heading: str, current_y: float):
            page, current_y = check_page_break(doc, page, current_y, 25)
            page.insert_text((margin_left, current_y), heading.upper(), fontsize=11, fontname="helv", color=(0.12, 0.35, 0.65))
            current_y += 5
            page.draw_line((margin_left, current_y), (margin_right, current_y), color=(0.85, 0.88, 0.92), width=0.5)
            current_y += 12
            return page, current_y

        # 2. Professional Summary
        summary = resume_data.get("summary")
        if summary:
            page, y = insert_section_heading(doc, page, "Professional Summary", y)
            # Wrap text roughly ~95 chars per line
            words = summary.split()
            current_line = []
            for w in words:
                current_line.append(w)
                if len(" ".join(current_line)) > 90:
                    page, y = check_page_break(doc, page, y, 14)
                    page.insert_text((margin_left, y), " ".join(current_line), fontsize=9.5, fontname="helv", color=(0.2, 0.2, 0.25))
                    y += 13
                    current_line = []
            if current_line:
                page, y = check_page_break(doc, page, y, 14)
                page.insert_text((margin_left, y), " ".join(current_line), fontsize=9.5, fontname="helv", color=(0.2, 0.2, 0.25))
                y += 16

        # 3. Technical Skills
        skills = resume_data.get("skills") or []
        if skills:
            page, y = insert_section_heading(doc, page, "Technical Skills", y)
            skills_text = ", ".join(skills)
            words = skills_text.split(", ")
            current_line = []
            for w in words:
                current_line.append(w)
                if len(", ".join(current_line)) > 85:
                    page, y = check_page_break(doc, page, y, 14)
                    page.insert_text((margin_left, y), ", ".join(current_line) + ",", fontsize=9.5, fontname="helv", color=(0.2, 0.2, 0.25))
                    y += 13
                    current_line = []
            if current_line:
                page, y = check_page_break(doc, page, y, 14)
                page.insert_text((margin_left, y), ", ".join(current_line), fontsize=9.5, fontname="helv", color=(0.2, 0.2, 0.25))
                y += 16

        # 4. Experience
        experience = resume_data.get("experience") or []
        if experience:
            page, y = insert_section_heading(doc, page, "Work Experience", y)
            for exp in experience:
                title = exp.get("title") or "Software Engineer"
                company = exp.get("company") or "Company"
                dates = f"{exp.get('startDate') or ''} - {'Present' if exp.get('current') else exp.get('endDate') or ''}".strip(" -")
                role_line = f"{title} | {company}"

                page, y = check_page_break(doc, page, y, 22)
                page.insert_text((margin_left, y), role_line, fontsize=10, fontname="helv", color=(0.1, 0.15, 0.2))
                if dates:
                    page.insert_text((margin_right - len(dates) * 6, y), dates, fontsize=9, fontname="helv", color=(0.4, 0.45, 0.5))
                y += 13

                bullets = exp.get("bullets") or []
                for b in bullets:
                    words = str(b).split()
                    current_line = ["•"]
                    for w in words:
                        current_line.append(w)
                        if len(" ".join(current_line)) > 88:
                            page, y = check_page_break(doc, page, y, 13)
                            page.insert_text((margin_left + 8, y), " ".join(current_line), fontsize=9, fontname="helv", color=(0.25, 0.25, 0.3))
                            y += 12
                            current_line = [" "]
                    if len(current_line) > 1:
                        page, y = check_page_break(doc, page, y, 13)
                        page.insert_text((margin_left + 8, y), " ".join(current_line), fontsize=9, fontname="helv", color=(0.25, 0.25, 0.3))
                        y += 12
                y += 5

        # 5. Projects
        projects = resume_data.get("projects") or []
        if projects:
            page, y = insert_section_heading(doc, page, "Projects", y)
            for proj in projects:
                p_name = proj.get("name") or "Project"
                techs = ", ".join(proj.get("technologies") or [])
                heading_line = f"{p_name} ({techs})" if techs else p_name

                page, y = check_page_break(doc, page, y, 18)
                page.insert_text((margin_left, y), heading_line, fontsize=9.5, fontname="helv", color=(0.15, 0.15, 0.2))
                y += 12

                desc = proj.get("description") or ""
                if desc:
                    words = desc.split()
                    current_line = ["•"]
                    for w in words:
                        current_line.append(w)
                        if len(" ".join(current_line)) > 88:
                            page, y = check_page_break(doc, page, y, 13)
                            page.insert_text((margin_left + 8, y), " ".join(current_line), fontsize=9, fontname="helv", color=(0.25, 0.25, 0.3))
                            y += 12
                            current_line = [" "]
                    if len(current_line) > 1:
                        page, y = check_page_break(doc, page, y, 13)
                        page.insert_text((margin_left + 8, y), " ".join(current_line), fontsize=9, fontname="helv", color=(0.25, 0.25, 0.3))
                        y += 12
                y += 4

        # 6. Education
        education = resume_data.get("education") or []
        if education:
            page, y = insert_section_heading(doc, page, "Education", y)
            for edu in education:
                inst = edu.get("institution") or ""
                deg = edu.get("degree") or ""
                field = edu.get("field") or ""
                deg_field = f"{deg} in {field}" if deg and field else deg or field
                edu_line = f"{deg_field} - {inst}" if deg_field else inst

                page, y = check_page_break(doc, page, y, 16)
                page.insert_text((margin_left, y), edu_line, fontsize=9.5, fontname="helv", color=(0.2, 0.2, 0.25))
                y += 14

        pdf_bytes = doc.tobytes()
        doc.close()
        return pdf_bytes

    @staticmethod
    def export_resume_docx(resume_data: Dict[str, Any]) -> bytes:
        """Generate an ATS-friendly, clean .docx document using python-docx."""
        doc = docx.Document()

        # Set 0.75 in margins
        for section in doc.sections:
            section.top_margin = Inches(0.75)
            section.bottom_margin = Inches(0.75)
            section.left_margin = Inches(0.75)
            section.right_margin = Inches(0.75)

        header = resume_data.get("header") or {}
        name = header.get("name") or "Candidate Name"
        contact_items = [
            header.get("email"),
            header.get("phone"),
            header.get("location"),
            header.get("linkedin"),
            header.get("github")
        ]
        contact_line = " | ".join([c for c in contact_items if c])

        # Header Title
        title_p = doc.add_paragraph()
        title_run = title_p.add_run(name)
        title_run.font.size = Pt(18)
        title_run.font.bold = True
        title_run.font.name = "Arial"

        if contact_line:
            contact_p = doc.add_paragraph()
            contact_run = contact_p.add_run(contact_line)
            contact_run.font.size = Pt(9.5)
            contact_run.font.color.rgb = RGBColor(90, 100, 110)
            contact_run.font.name = "Arial"

        def add_heading(text: str):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(12)
            p.paragraph_format.space_after = Pt(4)
            run = p.add_run(text.upper())
            run.font.size = Pt(11)
            run.font.bold = True
            run.font.color.rgb = RGBColor(25, 75, 140)
            run.font.name = "Arial"

        # Summary
        summary = resume_data.get("summary")
        if summary:
            add_heading("Professional Summary")
            p = doc.add_paragraph()
            r = p.add_run(summary)
            r.font.size = Pt(10)
            r.font.name = "Arial"

        # Skills
        skills = resume_data.get("skills") or []
        if skills:
            add_heading("Technical Skills")
            p = doc.add_paragraph()
            r = p.add_run(", ".join(skills))
            r.font.size = Pt(10)
            r.font.name = "Arial"

        # Experience
        experience = resume_data.get("experience") or []
        if experience:
            add_heading("Work Experience")
            for exp in experience:
                title = exp.get("title") or "Software Engineer"
                company = exp.get("company") or "Company"
                dates = f"{exp.get('startDate') or ''} - {'Present' if exp.get('current') else exp.get('endDate') or ''}".strip(" -")

                role_p = doc.add_paragraph()
                role_p.paragraph_format.space_before = Pt(6)
                role_p.paragraph_format.space_after = Pt(2)
                r_title = role_p.add_run(f"{title} | {company}")
                r_title.font.bold = True
                r_title.font.size = Pt(10.5)
                r_title.font.name = "Arial"

                if dates:
                    r_dates = role_p.add_run(f"   ({dates})")
                    r_dates.font.size = Pt(9.5)
                    r_dates.font.color.rgb = RGBColor(100, 110, 120)
                    r_dates.font.name = "Arial"

                bullets = exp.get("bullets") or []
                for b in bullets:
                    bp = doc.add_paragraph(style="List Bullet")
                    bp.paragraph_format.space_after = Pt(2)
                    br = bp.add_run(str(b))
                    br.font.size = Pt(9.5)
                    br.font.name = "Arial"

        # Projects
        projects = resume_data.get("projects") or []
        if projects:
            add_heading("Projects")
            for proj in projects:
                p_name = proj.get("name") or "Project"
                techs = ", ".join(proj.get("technologies") or [])
                head_text = f"{p_name} ({techs})" if techs else p_name

                pp = doc.add_paragraph()
                pp.paragraph_format.space_before = Pt(4)
                pp.paragraph_format.space_after = Pt(2)
                pr = pp.add_run(head_text)
                pr.font.bold = True
                pr.font.size = Pt(10)
                pr.font.name = "Arial"

                desc = proj.get("description") or ""
                if desc:
                    dp = doc.add_paragraph(style="List Bullet")
                    dp.paragraph_format.space_after = Pt(2)
                    dr = dp.add_run(desc)
                    dr.font.size = Pt(9.5)
                    dr.font.name = "Arial"

        # Education
        education = resume_data.get("education") or []
        if education:
            add_heading("Education")
            for edu in education:
                inst = edu.get("institution") or ""
                deg = edu.get("degree") or ""
                field = edu.get("field") or ""
                deg_field = f"{deg} in {field}" if deg and field else deg or field
                edu_line = f"{deg_field} - {inst}" if deg_field else inst

                ep = doc.add_paragraph()
                ep.paragraph_format.space_after = Pt(2)
                er = ep.add_run(edu_line)
                er.font.size = Pt(10)
                er.font.name = "Arial"

        stream = io.BytesIO()
        doc.save(stream)
        return stream.getvalue()

    @staticmethod
    def export_cover_letter_pdf(title: str, content: str, candidate_name: str = "") -> bytes:
        doc = fitz.open()
        page = doc.new_page(width=612, height=792)
        margin_left = 54.0
        y = 60.0

        if candidate_name:
            page.insert_text((margin_left, y), candidate_name, fontsize=16, fontname="helv", color=(0.1, 0.1, 0.15))
            y += 20

        page.insert_text((margin_left, y), title, fontsize=12, fontname="helv", color=(0.2, 0.35, 0.6))
        y += 10
        page.draw_line((margin_left, y), (558.0, y), color=(0.8, 0.8, 0.85), width=0.75)
        y += 24

        paragraphs = content.split("\n\n")
        for para in paragraphs:
            lines = para.strip().split("\n")
            for l in lines:
                words = l.split()
                current_line = []
                for w in words:
                    current_line.append(w)
                    if len(" ".join(current_line)) > 85:
                        page.insert_text((margin_left, y), " ".join(current_line), fontsize=10, fontname="helv", color=(0.2, 0.2, 0.25))
                        y += 14
                        current_line = []
                if current_line:
                    page.insert_text((margin_left, y), " ".join(current_line), fontsize=10, fontname="helv", color=(0.2, 0.2, 0.25))
                    y += 14
            y += 12

        pdf_bytes = doc.tobytes()
        doc.close()
        return pdf_bytes

    @staticmethod
    def export_cover_letter_docx(title: str, content: str, candidate_name: str = "") -> bytes:
        doc = docx.Document()
        for section in doc.sections:
            section.top_margin = Inches(1.0)
            section.bottom_margin = Inches(1.0)
            section.left_margin = Inches(1.0)
            section.right_margin = Inches(1.0)

        if candidate_name:
            hp = doc.add_paragraph()
            hr = hp.add_run(candidate_name)
            hr.font.size = Pt(16)
            hr.font.bold = True
            hr.font.name = "Arial"

        tp = doc.add_paragraph()
        tr = tp.add_run(title)
        tr.font.size = Pt(12)
        tr.font.bold = True
        tr.font.color.rgb = RGBColor(30, 80, 150)
        tr.font.name = "Arial"
        tp.paragraph_format.space_after = Pt(18)

        paragraphs = content.split("\n\n")
        for para in paragraphs:
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(12)
            p.paragraph_format.line_spacing = 1.15
            r = p.add_run(para.strip())
            r.font.size = Pt(11)
            r.font.name = "Arial"

        stream = io.BytesIO()
        doc.save(stream)
        return stream.getvalue()

document_exporter = DocumentExporter()
